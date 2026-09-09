"""
Exporta as tabelas do sistema_contabil.db para um arquivo Excel (.xlsx),
com uma aba por tabela/relatório.

Como usar:
    python exportar_excel.py

Gera o arquivo: relatorio_sistema_contabil.xlsx na mesma pasta.
Rode de novo sempre que quiser atualizar os dados no Excel.
"""
import sqlite3
import openpyxl
from openpyxl.utils import get_column_letter

DB_PATH = "sistema_contabil.db"
SAIDA = "relatorio_sistema_contabil.xlsx"

TABELAS_E_VIEWS = [
    ("Lançamentos", """
        SELECT l.id, e.razao_social AS empresa, l.data_lancamento, l.tipo,
               c.nome AS categoria, l.descricao, l.valor, l.status
        FROM lancamentos l
        JOIN empresas e ON e.id = l.empresa_id
        JOIN categorias c ON c.id = l.categoria_id
        ORDER BY l.data_lancamento DESC
    """),
    ("Vendas e Serviços", """
        SELECT v.id, e.razao_social AS empresa, v.data_venda, ps.tipo,
               ps.nome AS item, v.quantidade, v.valor_unitario, v.valor_total,
               v.cliente_nome, v.status
        FROM vendas v
        JOIN empresas e ON e.id = v.empresa_id
        JOIN produtos_servicos ps ON ps.id = v.produto_servico_id
        ORDER BY v.data_venda DESC
    """),
    ("Catálogo (Produtos-Serviços)", """
        SELECT ps.id, e.razao_social AS empresa, ps.codigo, ps.tipo, ps.nome,
               ps.unidade, ps.preco, ps.estoque_atual, ps.ativo
        FROM produtos_servicos ps
        JOIN empresas e ON e.id = ps.empresa_id
        ORDER BY e.razao_social, ps.nome
    """),
    ("Entradas de Mercadoria", """
        SELECT en.id, e.razao_social AS empresa, en.data_entrada, ps.codigo,
               ps.nome AS item, en.quantidade, en.valor_unitario, en.valor_total,
               en.fornecedor, en.status
        FROM entradas_estoque en
        JOIN empresas e ON e.id = en.empresa_id
        JOIN produtos_servicos ps ON ps.id = en.produto_servico_id
        ORDER BY en.data_entrada DESC
    """),
    ("Posição de Estoque", """
        SELECT e.razao_social AS empresa, ps.codigo, ps.nome, ps.unidade,
               ps.estoque_atual, ps.preco AS preco_venda,
               (ps.estoque_atual * ps.preco) AS valor_em_estoque
        FROM produtos_servicos ps
        JOIN empresas e ON e.id = ps.empresa_id
        WHERE ps.tipo = 'produto' AND ps.ativo = 1
        ORDER BY e.razao_social, ps.nome
    """),
    ("Fluxo de Caixa Mensal", "SELECT * FROM vw_fluxo_caixa_mensal ORDER BY empresa_id, mes"),
    ("DRE por Categoria", "SELECT * FROM vw_totais_por_categoria ORDER BY empresa_id, mes"),
    ("Empresas", "SELECT id, razao_social, nome_fantasia, cnpj, ativo FROM empresas"),
]

def exportar():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove a aba em branco padrão

    for nome_aba, query in TABELAS_E_VIEWS:
        cursor = conn.execute(query)
        colunas = [d[0] for d in cursor.description]
        linhas = cursor.fetchall()

        aba = wb.create_sheet(title=nome_aba[:31])  # Excel limita nome de aba a 31 chars
        aba.append(colunas)
        for linha in linhas:
            aba.append(list(linha))

        # deixa as colunas com largura razoável
        for i, col in enumerate(colunas, start=1):
            aba.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 2)

        # negrito no cabeçalho
        for cel in aba[1]:
            cel.font = openpyxl.styles.Font(bold=True)

    conn.close()
    wb.save(SAIDA)
    print(f"Exportado com sucesso: {SAIDA}")

if __name__ == "__main__":
    exportar()
